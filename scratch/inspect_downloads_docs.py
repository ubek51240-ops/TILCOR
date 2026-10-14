import os
import sys
import docx
import pymupdf

sys.stdout.reconfigure(encoding='utf-8')

# 1. Inspect docx
docx_path = r'C:\Users\user\Downloads\Telegram Desktop\Dissertatsiya_PhD_Kollokatsiya_Parallel_Korpus.docx'
if os.path.exists(docx_path):
    print(f"=== READING DOCX: {docx_path} ===")
    doc = docx.Document(docx_path)
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    print(f"Total non-empty paragraphs: {len(paras)}")
    print(f"Tables count: {len(doc.tables)}")
    # check for 535
    for i, p in enumerate(paras):
        if '535' in p or 'kecha' in p.lower():
            print(f"Para {i}: {p[:120]}")
    for t_idx, table in enumerate(doc.tables):
        print(f"Table {t_idx}: {len(table.rows)} rows, {len(table.columns)} cols")
        # Sample first row
        cells = [c.text.strip().replace('\n', ' ') for c in table.rows[0].cells]
        print(f"   Headers: {cells[:5]}")
        if len(table.rows) > 1:
            cells1 = [c.text.strip().replace('\n', ' ') for c in table.rows[1].cells]
            print(f"   Row 1: {cells1[:5]}")

# 2. Inspect pdf
pdf_path = r'C:\Users\user\Downloads\Telegram Desktop\Antigravity_Parallel_Corpus_PhD_Master_Prompt_UZ.pdf'
if os.path.exists(pdf_path):
    print(f"\n=== READING PDF: {pdf_path} ===")
    pdf = pymupdf.open(pdf_path)
    print(f"Pages: {len(pdf)}")
    for pno in range(len(pdf)):
        txt = pdf[pno].get_text()
        if '535' in txt or 'kecha' in txt.lower():
            for line in txt.splitlines():
                if '535' in line or 'kecha va kunduz' in line.lower():
                    print(f"Page {pno+1}: {line.strip()[:120]}")
